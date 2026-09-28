"""PPTX slide parser: text + notes + size checks data."""
from pathlib import Path
from typing import Dict, Any, List
from src.utils.text_clean import clean_text


def parse_pptx(path: str) -> Dict[str, Any]:
    p = Path(path)
    if not p.exists():
        return {"ok": False, "error": "file not found", "path": path}
    try:
        from pptx import Presentation
    except Exception as e:
        return {"ok": False, "error": f"python-pptx missing: {e}", "path": path}
    try:
        prs = Presentation(str(p))
        slides: List[Dict[str, Any]] = []
        for idx, slide in enumerate(prs.slides, start=1):
            texts: List[str] = []
            title = ""
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                t = clean_text(shape.text)
                if not t:
                    continue
                texts.append(t)
                if not title and shape.text_frame.paragraphs:
                    title = t.split("\n")[0][:120]
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = clean_text(slide.notes_slide.notes_text_frame.text)
            slides.append(
                {
                    "n": idx,
                    "title": title,
                    "bullets": texts,
                    "notes": notes,
                    "notes_len": len(notes),
                }
            )
        return {"ok": True, "path": str(p), "slides": len(slides), "items": slides}
    except Exception as e:
        return {"ok": False, "error": str(e), "path": str(p)}
